import os, sys, time, threading, json, traceback, logging, hmac
from logging.handlers import RotatingFileHandler
from datetime import datetime

# ── Persistent file logging (5MB × 3 backups) ───────────────
_log_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'bot.log')
_file_handler = RotatingFileHandler(_log_path, maxBytes=5_000_000, backupCount=3, encoding='utf-8')
_file_handler.setFormatter(logging.Formatter('%(asctime)s [%(levelname)s] %(message)s', datefmt='%Y-%m-%d %H:%M:%S'))
logging.getLogger().addHandler(_file_handler)
logging.getLogger().setLevel(logging.INFO)
logging.getLogger('httpx').setLevel(logging.WARNING)
logging.getLogger('httpcore').setLevel(logging.WARNING)
logging.getLogger('werkzeug').setLevel(logging.WARNING)
_botlog = logging.getLogger('bot')

# Windows: ép UTF-8 cho stdout/stderr để print tiếng Việt không crash khi redirect
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

import openpyxl
from flask import Flask, jsonify, render_template, Response, stream_with_context, request, abort

from strategy import (
    get_funding_rates, get_available_usdt, get_balance_snapshot, get_funding_income,
    open_position, close_position, sell_spot,
    check_exit_conditions, estimate_pnl, get_spot_price,
    get_okx_swap_positions, get_instrument_state, funding_payments_since, collectible_rate,
    get_all_spot_balances, get_swap_info,
    MIN_FUNDING_RATE, MIN_USDT, PRICE_STOP_PCT,
    ROUND_TRIP_FEE, FEE_SAFETY, ENTRY_MIN_SETTLEMENTS, EXIT_MIN_SETTLEMENTS,
    adaptive_position_pct, SCAN_COINS, validate_scan_coins,
    CAPITAL_FRACTION, CAPITAL_PER_NOTIONAL,
)
from cross_bot_lock import account_lock
from excel_logger import log_pnl_snapshot, export_json, push_to_github, EXCEL_FILE
import analytics
import notifier
from config import SIMULATED
from ws_ticker import TickerWS

analytics.init()
analytics.backfill_fees()  # cập nhật fee cho trade cũ trong DB

# WebSocket realtime — subscribe SPOT tickers cho mọi coin trong watchlist
_ws = TickerWS([f"{c}-USDT" for c in SCAN_COINS], simulated=SIMULATED, max_age=15)


def _get_spot_price(spot_id):
    """Ưu tiên giá từ WS cache (sub-second), fallback REST."""
    p = _ws.get_price(spot_id)
    if p is not None:
        return p
    return get_spot_price(spot_id)

app = Flask(__name__)

# ── Bảo vệ endpoint đổi-trạng-thái (audit critical: /api/* POST không auth) ──
# Mặc định bind 127.0.0.1 → vào qua SSH tunnel (request đến như loopback, tin cậy).
# Nếu BIND_HOST=0.0.0.0 (public): mọi POST KHÔNG-loopback bắt buộc header X-Dash-Token == DASH_TOKEN.
_DASH_TOKEN = os.getenv('DASH_TOKEN', '').strip()

@app.before_request
def _guard_mutations():
    if request.method in ('POST', 'PUT', 'DELETE'):
        if request.remote_addr in ('127.0.0.1', '::1'):
            return
        sent = request.headers.get('X-Dash-Token', '')
        if not _DASH_TOKEN or not hmac.compare_digest(sent, _DASH_TOKEN):
            abort(401)

_lock  = threading.Lock()
_state = {
    'running':       False,
    'positions':     [],
    'opportunities': [],
    'usdt':          0.0,
    'total_eq':      0.0,     # tổng tài sản (totalEq USD) — gồm spot chân long + margin swap
    'log':           [],
    'last_update':   '-',
    'closing':       set(),   # coin đang trong tiến trình đóng (chống double-close)
    'io_busy':       set(),   # job IO đang chạy (excel, push) — chống stacking
}

MAX_POS         = 10       # 4→10: nhiều slot delta-neutral hơn để xoay vòng theo funding ⇒ tiến gần
                           # 20 lệnh/ngày (3 settlement × rotate). Không thêm rủi ro HƯỚNG GIÁ (đã hedge),
                           # chỉ dùng thêm margin — đã hạ POSITION_PCT 15→7% cho khớp.
SCAN_INT        = 180       # scan funding rate mỗi 3 phút (300→180: lấp slot vừa giải phóng nhanh hơn)
MON_INT         = 2         # đọc giá WS + cập nhật PnL MỖI 2 GIÂY (realtime)
BAL_INT         = 15        # đọc số dư REST mỗi 15s (tách khỏi MON 2s — đỡ đập API)
FUNDING_CHK_INT = 30        # check funding rate exit-condition mỗi 30s (rate đổi 8h/lần)
TICK_LOG_INT    = 60        # ghi tick log vào DB mỗi 60s
RECON_INT       = 45        # reconcile với OKX mỗi 45s (150→45: phát hiện phantom & retry dọn
                            # spot unhedged nhanh hơn — quan trọng khi chung TK với trend-bot)
ORPHAN_INT      = 120       # quét spot mồ côi (đã mua spot nhưng KHÔNG có swap hedge) mỗi 2 phút.
                            # Chạy ĐỘC LẬP với reconcile thường: orphan có thể tồn tại khi KHÔNG
                            # có vị thế local nào (vd crash giữa 2 chân mở lệnh).
WAL_CKPT_INT    = 3600      # checkpoint WAL DB mỗi giờ (chống file .db-wal phình vô hạn)
EXCL_INT        = 8*3600    # ghi Excel mỗi 8 giờ
PUSH_INT        = 300       # push GitHub Pages mỗi 5 phút
TICK            = 1         # vòng lặp chính 1s (UI cảm nhận realtime)
MAX_LOG         = 300
SSE_TTL         = 1800
_bot_thread = None

STATE_FILE = os.path.join(os.path.dirname(__file__), 'state.json')


# ── Persistence ──────────────────────────────────────────────────
def _persist_state():
    """Lưu vị thế đang mở ra disk (atomic write)."""
    try:
        with _lock:
            snap = [
                {k: v for k, v in p.items() if not k.startswith('_') or k == '_trade_id'}
                for p in _state['positions']
            ]
        tmp = STATE_FILE + '.tmp'
        with open(tmp, 'w', encoding='utf-8') as f:
            json.dump({'positions': snap, 'ts': time.time()}, f, ensure_ascii=False)
        os.replace(tmp, STATE_FILE)
    except Exception as e:
        # Không cho lỗi persist phá vỡ bot — chỉ log
        try: _log(f"Lưu state.json lỗi: {e}")
        except Exception: pass


def _load_state():
    if not os.path.exists(STATE_FILE):
        return []
    try:
        with open(STATE_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
        return data.get('positions', []) or []
    except Exception:
        return []


# ── Race-condition guard ─────────────────────────────────────────
def _begin_close(coin):
    """Trả True nếu đoạt được quyền đóng coin này; False nếu đã có thread khác."""
    with _lock:
        if coin in _state['closing']:
            return False
        _state['closing'].add(coin)
    return True


def _end_close(coin):
    with _lock:
        _state['closing'].discard(coin)


def _arm_close_backoff(p):
    """Đóng thất bại → lùi lịch thử lại (exponential, tối đa 60s) thay vì đập API mỗi 2s."""
    n = p.get('_close_attempts', 0)
    p['_close_attempts']   = n + 1
    p['_close_fail_until'] = time.time() + min(60, 5 * (2 ** n))
    p['_exit'] = False  # sẽ tự bật lại ở chu kỳ sau nếu điều kiện thoát vẫn đúng


# ── IO offload (fire-and-forget worker để bot không bị treo bởi git/Excel) ─
def _offload(key, fn, *args):
    with _lock:
        if key in _state['io_busy']:
            return  # job cùng key đang chạy, bỏ qua chu kỳ này
        _state['io_busy'].add(key)
    def runner():
        try:
            fn(*args)
        except Exception as e:
            _log(f"Lỗi IO[{key}]: {e}")
        finally:
            with _lock:
                _state['io_busy'].discard(key)
    threading.Thread(target=runner, daemon=True, name=f"io-{key}").start()


def _log(msg: str):
    line = f"[{datetime.now().strftime('%H:%M:%S')}] {msg}"
    with _lock:
        _state['log'].append(line)
        if len(_state['log']) > MAX_LOG:
            _state['log'] = _state['log'][-MAX_LOG:]
    try: _botlog.info(msg)
    except Exception: pass


def _finalize_close(p, final_pnl, reason):
    """Ghi sổ + xóa local + notify. CHỈ gọi khi đã xác nhận futures hết VÀ spot đã sạch."""
    coin = p['coin']
    p['_close_attempts'] = 0
    p['_pending_spot_cleanup'] = False
    analytics.record_close(p, final_pnl, reason=reason)
    with _lock:
        _state['positions'] = [x for x in _state['positions'] if x['coin'] != coin]
    _persist_state()
    pnl_val = (final_pnl or {}).get('net_pnl') or (final_pnl or {}).get('total_pnl') or 0
    _log(f"[{coin}] Đóng ✓ ({reason})  net {pnl_val:+.2f}$")
    ev = 'CLOSE_WIN' if pnl_val >= 0 else 'CLOSE_LOSS'
    notifier.notify_event(ev, f"Đóng <b>{coin}</b> ({reason})  net {pnl_val:+.2f}$")


def _close_one(p, reason):
    """Đóng 1 vị thế arb. BẤT BIẾN AN TOÀN: chỉ xóa local khi futures ĐÃ hết VÀ spot
    ĐÃ bán sạch. Nếu spot chưa bán được → GIỮ local + cờ _pending_spot_cleanup +
    alert CRITICAL + retry (qua reconcile / monitor). KHÔNG để lại unhedged untracked."""
    coin = p['coin']
    if not _begin_close(coin):
        return False
    try:
        # PnL tươi để ghi analytics — funding lấy số THẬT từ bills, fallback cache nếu REST fail
        try:
            fa = get_funding_income(p['swap_id'], p['open_time'])
            if fa is None:
                fa = p.get('_funding_actual')
            final_pnl = estimate_pnl(p, funding_override=fa)
        except Exception:
            final_pnl = None
        final_pnl = final_pnl or p.get('_pnl')

        # ── Trường hợp futures ĐÃ đóng từ trước, chỉ còn dọn nốt spot ──
        # (KHÔNG gọi close_position → tránh reduceOnly-buy lên futures đã hết = 51169 storm)
        if p.get('_pending_spot_cleanup'):
            if sell_spot(p['spot_id'], p['coin_amount']):
                _finalize_close(p, final_pnl, reason)
                return True
            _arm_close_backoff(p)
            notifier.notify_critical(
                f"{coin}: futures đã đóng nhưng SPOT vẫn CHƯA bán được (UNHEDGED) — bot đang retry.",
                key=f"unhedged-{coin}")
            return False

        # ── Đóng bình thường: futures + spot ──
        try:
            close_ok = close_position(p)
        except Exception as e:
            _log(f"[{coin}] Exception khi close_position: {e}")
            close_ok = False

        # VERIFY futures với OKX trước khi xóa local
        time.sleep(1.0)  # cho exchange settle
        okx_pos = get_okx_swap_positions()
        if okx_pos is None:
            _log(f"[{coin}] ⚠ Không query được vị thế OKX khi đóng — backoff, thử lại sau")
            _arm_close_backoff(p)
            return False
        if coin in okx_pos:
            # Futures CHƯA đóng được. Phân biệt market đóng/tạm dừng vs lỗi tạm thời.
            # KHÔNG bao giờ báo 'đã đóng' khi futures còn trên sàn.
            state    = get_instrument_state(p['swap_id'])
            attempts = p.get('_close_attempts', 0)
            err      = p.get('_last_close_err')
            if state and state != 'live':
                _log(f"[{coin}] ⚠ Market '{state}' (không giao dịch) — KHÔNG đóng được futures, giữ + tự retry khi mở lại")
                notifier.notify_critical(
                    f"{coin}: KHÔNG đóng được futures vì instrument đang '{state}' (market đóng/tạm dừng). "
                    f"Vị thế short VẪN CÒN trên sàn — bot tự retry khi market mở lại. "
                    f"Cần đóng gấp thì xử lý thủ công trên OKX." + (f" Lý do sàn: {err}" if err else ""),
                    key=f"market-closed-{coin}")
            elif attempts >= 2:
                _log(f"[{coin}] ⚠ Đóng futures {attempts+1} lần CHƯA khớp (market 'live') — vẫn còn, đang retry")
                notifier.notify_critical(
                    f"{coin}: lệnh đóng futures đã thử {attempts+1} lần nhưng VẪN còn trên sàn (market 'live'). "
                    f"Bot tiếp tục retry — kiểm tra thủ công nếu kéo dài."
                    + (f" Lý do sàn: {err}" if err else ""),
                    key=f"close-stuck-{coin}")
            else:
                _log(f"[{coin}] ⚠ Sau khi đóng, OKX vẫn còn futures — backoff, thử lại sau")
            _arm_close_backoff(p)
            return False

        # Futures đã KHÔNG còn. BẮT BUỘC spot sạch trước khi xóa local.
        # close_position trả True = spot cũng đã bán; nếu False thì thử bán riêng (KHÔNG gọi lại close_position).
        spot_clean = bool(close_ok) or sell_spot(p['spot_id'], p['coin_amount'])
        if not spot_clean:
            p['_pending_spot_cleanup'] = True   # GIỮ local để retry, KHÔNG xóa
            _arm_close_backoff(p)
            _persist_state()
            notifier.notify_critical(
                f"{coin}: futures đã đóng nhưng SPOT CHƯA bán được → UNHEDGED. Bot sẽ tự retry bán spot.",
                key=f"unhedged-{coin}")
            return False

        _finalize_close(p, final_pnl, reason)
        return True
    except Exception as e:
        _log(f"[{coin}] Lỗi đóng: {e}")
        _arm_close_backoff(p)
        return False
    finally:
        _end_close(coin)


def _reconcile_remove(p):
    """Xóa vị thế khỏi local state vì futures đã không còn trên OKX (liquidated / đóng tay).
    Bán spot để tránh để lại unhedged position trên sàn."""
    coin = p['coin']
    if not _begin_close(coin):
        return 0
    try:
        try:
            fa = get_funding_income(p['swap_id'], p['open_time'])
            if fa is None:
                fa = p.get('_funding_actual')
            final_pnl = estimate_pnl(p, funding_override=fa) or {}
        except Exception:
            final_pnl = {}
        # Futures đã đóng bên ngoài — PHẢI bán nốt spot để hết unhedged.
        # Nếu bán fail → GIỮ local + cờ pending + alert; reconcile lần sau retry (KHÔNG bỏ rơi spot).
        if not sell_spot(p['spot_id'], p['coin_amount']):
            p['_pending_spot_cleanup'] = True
            _arm_close_backoff(p)
            _persist_state()
            notifier.notify_critical(
                f"{coin}: futures đóng bên ngoài nhưng SPOT chưa bán được (UNHEDGED) — bot sẽ retry.",
                key=f"unhedged-{coin}")
            return 0
        _finalize_close(p, final_pnl, reason='external_close')
        _log(f"[{coin}] ⚠ Reconciled — không còn trên OKX, đã đóng local + bán spot")
        return 1
    except Exception as e:
        _log(f"[{coin}] Lỗi reconcile: {e}")
        return 0
    finally:
        _end_close(coin)


def _reconcile_with_okx():
    """Query OKX, drop những local position không còn tồn tại. Trả số đã reconcile."""
    okx_pos = get_okx_swap_positions()
    if okx_pos is None:
        return 0  # API fail — không kết luận
    with _lock:
        local = list(_state['positions'])
    removed = 0
    for p in local:
        if p['coin'] not in okx_pos:
            removed += _reconcile_remove(p)
    return removed


def _reconcile_orphan_spot():
    """Quét số dư spot của các coin BOT GIAO DỊCH mà KHÔNG có chân swap hedge nào (local lẫn
    trên OKX) ⇒ 'spot mồ côi'. Gốc: bot bị kill ngay giữa lúc open_position đã MUA spot nhưng
    CHƯA kịp short swap (hoặc chưa kịp append vào state) → còn spot trần, KHÔNG có record local
    nên reconcile thường (chỉ duyệt local) không bao giờ thấy. Để lâu = rủi ro hướng giá không hedge.

    AN TOÀN:
      - CHỈ đụng coin trong SCAN_COINS (coin bot tự giao dịch) — không chạm số dư lạ.
      - Bỏ qua coin đang có swap (đã hedge), có vị thế local, hoặc đang trong tiến trình đóng.
      - Bỏ qua dust (< nửa MIN_USDT) để tránh bán nhầm vụn phí.
      - open_position chạy ĐỒNG BỘ trong cùng thread _bot() ⇒ không bao giờ interleave với quét này.
    Trả số coin đã dọn."""
    balances = get_all_spot_balances()
    if not balances:
        return 0
    okx_pos = get_okx_swap_positions()
    if okx_pos is None:
        return 0  # API fail — KHÔNG kết luận (tránh bán nhầm khi không đọc được swap)
    with _lock:
        local_coins = {p['coin'] for p in _state['positions']}
        closing     = set(_state['closing'])
    cleaned = 0
    for ccy, avail in balances.items():
        if ccy not in SCAN_COINS:
            continue
        if ccy in local_coins or ccy in okx_pos or ccy in closing:
            continue  # đã có chân swap hedge / đang xử lý → KHÔNG phải mồ côi
        spot_id = f"{ccy}-USDT"
        price = _get_spot_price(spot_id)
        if price is None:
            continue
        notional = avail * price
        if notional < MIN_USDT * 0.5:
            continue  # dust — bỏ qua
        _log(f"[{ccy}] ⚠ SPOT MỒ CÔI ${notional:.2f} (không có swap hedge) — dọn bán...")
        notifier.notify_critical(
            f"{ccy}: phát hiện spot KHÔNG hedge ${notional:.2f} (nghi bot crash giữa 2 chân mở lệnh) "
            f"— bot đang tự bán dọn.", key=f"orphan-{ccy}")
        if sell_spot(spot_id, avail):
            cleaned += 1
            _log(f"[{ccy}] Đã dọn spot mồ côi ✓")
        else:
            _log(f"[{ccy}] Bán spot mồ côi CHƯA được — sẽ thử lại chu kỳ sau")
    return cleaned


def _reconcile_adopt_untracked():
    """FIX (audit 07/2026): 'nhận lại' vị thế ĐÃ HEDGE (swap SHORT + spot khớp) tồn tại trên OKX
    nhưng KHÔNG có record local — xảy ra khi bot bị kill đúng khe giữa 'open_position xong' và
    'append vào state + persist'. Không adopt thì vị thế này bị bỏ quản lý VĨNH VIỄN (không bao giờ
    đóng theo funding_flip/price_stop, kẹt margin), vì reconcile thường chỉ duyệt local.

    CHỈ adopt khi CẢ 3 đúng: coin ∈ SCAN_COINS, có swap SHORT (pos<0) trên OKX, VÀ có spot balance
    khớp (≥50% notional swap ⇒ đã hedge). Không thoả (vd spot thiếu) → để _reconcile_orphan_spot xử lý.
    Trường entry (open_time/entry_rate) tái dựng best-effort (open_time=now → funding tính bảo thủ,
    KHÔNG over-report). Trả số vị thế đã adopt."""
    okx_pos = get_okx_swap_positions()
    if okx_pos is None:
        return 0
    balances = get_all_spot_balances()
    if balances is None:
        return 0
    with _lock:
        local_coins = {p['coin'] for p in _state['positions']}
        closing     = set(_state['closing'])
    adopted = 0
    for coin, info in okx_pos.items():
        if coin not in SCAN_COINS or coin in local_coins or coin in closing:
            continue
        if info['pos'] >= 0:
            continue  # arb luôn SHORT swap; pos≥0 không phải vị thế của bot
        swap_id = info['inst_id']
        ct_val, min_sz, lot_sz = get_swap_info(swap_id)
        if ct_val is None or not lot_sz:
            continue
        contracts  = abs(info['pos'])
        spot_avail = balances.get(coin, 0.0)
        if spot_avail < (contracts * ct_val) * 0.5:
            continue  # spot không đủ hedge → KHÔNG phải arb đã hedge (orphan/khác lo)
        entry_px = info['avg_px'] or _get_spot_price(f"{coin}-USDT")
        if not entry_px:
            continue
        # entry_funding_rate: lấy rate hiện tại làm fallback (PnL thật vẫn ưu tiên bills)
        try:
            _, cur_rate = check_exit_conditions({'swap_id': swap_id, 'coin': coin})
        except Exception:
            cur_rate = None
        pos = {
            'coin':               coin,
            'spot_id':            f"{coin}-USDT",
            'swap_id':            swap_id,
            'coin_amount':        round(spot_avail, 8),
            'contracts':          contracts,
            'lot_sz':             lot_sz,
            'ct_val':             ct_val,
            'entry_price':        entry_px,
            'entry_funding_rate': cur_rate if cur_rate is not None else 0.0001,
            'open_time':          time.time(),   # bảo thủ: không rõ mốc mở thật
            '_adopted':           True,
        }
        analytics.record_open(pos)   # tạo record để close sau này ghi sổ khớp
        with _lock:
            _state['positions'].append(pos)
        _persist_state()
        adopted += 1
        _log(f"[{coin}] ⚠ ADOPT vị thế hedge chưa tracked (swap {contracts:g} + spot {spot_avail:.6f}) — đưa vào quản lý")
        notifier.notify_critical(
            f"{coin}: phát hiện vị thế ĐÃ HEDGE trên sàn nhưng bot KHÔNG có record (nghi crash giữa "
            f"mở lệnh & ghi state). Bot đã nhận lại để quản lý/đóng bình thường.", key=f"adopt-{coin}")
    return adopted


def _bot():
    last_scan = last_mon = last_excel = last_push = last_recon = last_tick_log = last_fund_chk = 0
    last_bal = last_ckpt = last_orphan = 0
    last_opps: list = []
    _my_thread = threading.current_thread()

    _log("━━━ Bot OKX Funding Arb khởi động ━━━")

    # Lọc động coin đã delist/settle khỏi watchlist (vd TON 07/2026) TRƯỚC khi sub WS —
    # tránh sub instrument chết (ws 51001/60018) và tránh cố mở lệnh trên coin không còn 'live'.
    try:
        dropped = validate_scan_coins()
        if dropped:
            _log(f"⚠ Loại {len(dropped)} coin đã delist/không 'live' khỏi watchlist: {', '.join(dropped)}")
        _ws.instruments = [f"{c}-USDT" for c in SCAN_COINS]  # WS theo danh sách đã lọc
    except Exception as e:
        _log(f"⚠ validate_scan_coins lỗi (giữ nguyên watchlist): {e}")

    if _ws.start():
        _log(f"WebSocket realtime ticker → đang kết nối ({len(_ws.instruments)} sym)")

    # Restore vị thế đã lưu (nếu app từng dừng đột ngột)
    restored = _load_state()
    if restored:
        with _lock:
            _state['positions'] = restored
        _log(f"⚠ Restore {len(restored)} vị thế từ state.json — đang reconcile với OKX...")
        n = _reconcile_with_okx()
        if n:
            _log(f"Reconcile: xóa {n} vị thế phantom (không còn trên OKX)")

    # FIX (audit 07/2026): nhận lại vị thế đã hedge trên sàn nhưng mất record local (crash giữa
    # mở lệnh & ghi state) — chạy KỂ CẢ khi state.json trống.
    try:
        a = _reconcile_adopt_untracked()
        if a:
            _log(f"Adopt: nhận lại {a} vị thế hedge chưa tracked trên OKX")
    except Exception as e:
        _log(f"Lỗi adopt untracked lúc khởi động: {e}")

    _bal0, _teq0 = get_balance_snapshot()   # ngoài lock: tránh giữ lock khi gọi mạng
    with _lock:
        _state['usdt'] = _bal0
        if _teq0 > 0:
            _state['total_eq'] = _teq0
    _log(f"Số dư: ${_bal0:.2f} USDT khả dụng · tổng tài sản ${_teq0:.2f}")

    while True:
        with _lock:
            if not _state['running'] or _bot_thread is not _my_thread:
                break
        now = time.time()

        # ── Cập nhật PnL & kiểm tra thoát ───────────────────────
        if now - last_mon >= MON_INT:
            with _lock:
                positions = list(_state['positions'])

            do_tick_log  = (now - last_tick_log) >= TICK_LOG_INT
            do_fund_chk  = (now - last_fund_chk) >= FUNDING_CHK_INT
            for p in positions:
                try:
                    # ── TÍNH TOÁN / NETWORK (NGOÀI lock — không giữ lock khi gọi mạng) ──
                    cur_rate = None
                    fa_val   = None
                    exit_flag = False
                    exit_reason = None
                    # Funding rate + funding THỰC NHẬN chỉ check thưa (rate đổi mỗi 8h)
                    if do_fund_chk:
                        ok, cur_rate = check_exit_conditions(p)
                        fa = get_funding_income(p['swap_id'], p['open_time'])
                        if fa is not None:
                            fa_val = fa
                        if ok:
                            # EV-HONEST EXIT: chỉ thoát funding_flip khi funding ĐÃ THU đủ bù phí round-trip
                            # → mọi lệnh đóng ra luôn net-dương về funding. Hard-cap EXIT_MIN_SETTLEMENTS kỳ.
                            n_paid   = funding_payments_since(p['open_time'])
                            notional = p['contracts'] * p['ct_val'] * p['entry_price']
                            fee      = notional * ROUND_TRIP_FEE
                            got      = fa_val if fa_val is not None else p.get('_funding_actual')
                            if got is None:
                                got = (p.get('_pnl') or {}).get('funding_pnl', 0) or 0
                            if got >= fee or n_paid >= EXIT_MIN_SETTLEMENTS:
                                exit_flag, exit_reason = True, 'funding_flip'
                    # Giá WS (sub-second), fallback REST. Funding dùng số thật nếu có.
                    spot_price = _get_spot_price(p['spot_id'])
                    fund_override = fa_val if fa_val is not None else p.get('_funding_actual')
                    pnl = estimate_pnl(p, price=spot_price, funding_override=fund_override)
                    # Price stop — thoát nếu giá diverge quá PRICE_STOP_PCT
                    if not exit_flag and not p.get('_exit') and pnl:
                        notional = p['contracts'] * p['ct_val'] * p['entry_price']
                        if pnl.get('price_pnl', 0) < -(notional * PRICE_STOP_PCT):
                            exit_flag, exit_reason = True, 'price_stop'

                    # ── ÁP MUTATION VÀO vị thế (TRONG lock — chống race với persist/SSE/close) ──
                    with _lock:
                        if cur_rate is not None:
                            p['_cur_rate'] = cur_rate
                        if fa_val is not None:
                            p['_funding_actual'] = fa_val
                        p['_pnl'] = pnl
                        if exit_flag:
                            p['_exit'] = True
                            p['_exit_reason'] = exit_reason
                    if exit_reason == 'price_stop':
                        _log(f"[{p['coin']}] Price stop (-{PRICE_STOP_PCT*100:.0f}%) → đóng")
                    if do_tick_log and pnl:
                        analytics.record_tick(p, pnl, funding_rate=cur_rate)
                except Exception as e:
                    _log(f"[{p['coin']}] Lỗi cập nhật: {e}")
            if do_tick_log: last_tick_log = now
            if do_fund_chk: last_fund_chk = now

            now_close = time.time()
            # Đóng các vị thế cần thoát + retry các vị thế đang kẹt dọn spot (unhedged) — đều tôn trọng backoff.
            for p in [x for x in positions
                      if (x.get('_exit') or x.get('_pending_spot_cleanup'))
                      and now_close >= x.get('_close_fail_until', 0)]:
                reason = p.get('_exit_reason') or 'funding_flip'
                if p.get('_pending_spot_cleanup'):
                    _log(f"[{p['coin']}] Retry dọn spot (unhedged)...")
                else:
                    label = 'Price stop' if reason == 'price_stop' else 'Funding rate thấp'
                    _log(f"[{p['coin']}] {label} → đóng vị thế...")
                _close_one(p, reason=reason)

            # Số dư đọc thưa hơn (REST) — không cần realtime như giá
            if now - last_bal >= BAL_INT:
                bal, teq = get_balance_snapshot()   # ngoài lock: tránh giữ lock khi gọi mạng
                with _lock:
                    _state['usdt'] = bal
                    if teq > 0:
                        _state['total_eq'] = teq
                last_bal = now
            with _lock:
                _state['last_update'] = datetime.now().strftime('%H:%M:%S')
            last_mon = now

        # ── Reconcile với OKX (xóa phantom positions) ────────────
        if now - last_recon >= RECON_INT:
            with _lock:
                has_local = bool(_state['positions'])
            if has_local:
                try:
                    n = _reconcile_with_okx()
                    if n:
                        _log(f"Reconcile tự động: xóa {n} vị thế phantom")
                except Exception as e:
                    _log(f"Lỗi reconcile: {e}")
            last_recon = now

        # ── Quét spot mồ côi (chạy KỂ CẢ khi không có vị thế local) ──
        if now - last_orphan >= ORPHAN_INT:
            try:
                n = _reconcile_orphan_spot()
                if n:
                    _log(f"Dọn {n} spot mồ côi (mua spot nhưng thiếu swap hedge)")
            except Exception as e:
                _log(f"Lỗi quét spot mồ côi: {e}")
            # FIX (audit 07/2026): nhận lại vị thế hedge chưa tracked (cùng nhịp với orphan-scan)
            try:
                a = _reconcile_adopt_untracked()
                if a:
                    _log(f"Adopt: nhận lại {a} vị thế hedge chưa tracked trên OKX")
            except Exception as e:
                _log(f"Lỗi adopt untracked: {e}")
            last_orphan = now

        # ── Scan cơ hội ───────────────────────────────────────────
        if now - last_scan >= SCAN_INT:
            with _lock:
                n_pos = len(_state['positions'])

            if n_pos < MAX_POS:
                _log(f"SCAN — ${_state['usdt']:.2f} USDT  |  {datetime.now().strftime('%H:%M')}")
                opps = get_funding_rates()
                if opps:
                    analytics.record_scan(opps)
                    opps = analytics.rank_opps(opps)
                last_opps = opps or []
                with _lock:
                    _state['opportunities'] = last_opps[:10]

                if opps:
                    with _lock:
                        open_coins = {p['coin'] for p in _state['positions']}
                    # Snapshot vị thế trên OKX để biết coin nào của bot KHÁC (tránh đụng khi chung TK).
                    okx_now = get_okx_swap_positions() or {}
                    for opp in opps:
                        coin = opp['coin']
                        with _lock:
                            n_pos = len(_state['positions'])
                            go    = _state['running']
                            own   = {p['coin'] for p in _state['positions']}
                        if not go or n_pos >= MAX_POS:
                            break
                        if coin in open_coins:
                            continue
                        if opp['funding_rate'] < MIN_FUNDING_RATE:
                            break
                        # OWNERSHIP: coin đã có vị thế trên OKX nhưng KHÔNG thuộc bot này → của bot kia
                        # (chung tài khoản). Mở thêm sẽ NET vào nhau → reduceOnly đóng nhầm. Bỏ qua.
                        if coin in okx_now and coin not in own:
                            _log(f"[{coin}] Bỏ qua — đã có vị thế OKX của bot khác (chung TK), tránh đụng")
                            continue
                        # Cổng EV theo funding DỰ BÁO thực thu kỳ kế (next_rate).
                        coll = collectible_rate(opp)
                        if coll * ENTRY_MIN_SETTLEMENTS < ROUND_TRIP_FEE * FEE_SAFETY:
                            _log(f"[{coin}] Bỏ qua — EV thấp (thực thu {coll*100:.4f}%×{ENTRY_MIN_SETTLEMENTS}kỳ "
                                 f"< phí {ROUND_TRIP_FEE*FEE_SAFETY*100:.3f}% · cur={opp['funding_rate']*100:.3f}% next={opp['next_rate']*100:.3f}%)")
                            continue
                        if opp.get('hist_score', 0) < -1.5:
                            _log(f"[{coin}] Bỏ qua — lịch sử kém (score={opp['hist_score']:.2f})")
                            continue

                        # ── MUTEX liên-bot: nối tiếp việc đặt lệnh để 2 bot không tranh margin/spot ──
                        pos = None
                        with account_lock('arb-open') as ok:
                            if not ok:
                                _log("Không lấy được khóa giao dịch (bot kia đang đặt lệnh) — dừng scan, thử lượt sau")
                                break
                            # OWNERSHIP RE-CHECK ATOMIC DƯỚI KHÓA (fix net_mode chân short tan rã):
                            # snapshot okx_now ở đầu scan có thể CŨ — bot kia có thể vừa mở coin này.
                            # Vì CẢ 2 bot mở lệnh dưới CÙNG file-lock (tuần tự), re-đọc vị thế OKX lúc
                            # này CHẮC CHẮN thấy vị thế bot kia vừa mở ⇒ không bao giờ net cùng instrument.
                            okx_fresh = get_okx_swap_positions()
                            if okx_fresh is not None and coin in okx_fresh and coin not in own:
                                _log(f"[{coin}] Bỏ qua (re-check dưới khóa) — coin đã có vị thế bot khác (chung TK)")
                                continue
                            # Re-đọc số dư TRONG khóa cho tươi; tính ngân sách arb (capital fraction).
                            available = get_available_usdt()
                            with _lock:
                                deployed = sum(x['contracts'] * x['ct_val'] * x['entry_price']
                                               for x in _state['positions']) * CAPITAL_PER_NOTIONAL
                            total_equity = available + deployed
                            budget = total_equity * CAPITAL_FRACTION
                            room   = budget - deployed
                            max_notional = min(available, room) / CAPITAL_PER_NOTIONAL
                            if max_notional < MIN_USDT:
                                _log(f"Hết ngân sách arb (budget ${budget:.0f}, đã dùng ${deployed:.0f}/{CAPITAL_FRACTION*100:.0f}% TK) — dừng scan")
                                break
                            amount = min(total_equity * adaptive_position_pct(opp['funding_rate']), max_notional)
                            if amount < MIN_USDT:
                                break
                            hs = opp.get('hist_score', 0)
                            hs_tag = f" · score {hs:+.2f}" if hs else ''
                            _log(f"[{coin}] Vào lệnh ${amount:.2f} @ {opp['funding_rate']*100:.4f}%/8h{hs_tag}")
                            # FIX (audit 07/2026): bọc try/except — open_position có write-path gọi
                            # API trực tiếp; nếu ném exception thì open_position tự rollback spot, nhưng
                            # exception KHÔNG được để lan ra đá cả _bot() loop về restart (backoff 30-300s).
                            try:
                                pos = open_position(opp, amount)
                            except Exception as e:
                                _log(f"[{coin}] ⚠ Lỗi mở lệnh (đã bỏ qua, orphan-scan sẽ dọn spot nếu sót): {e}")
                                pos = None
                        # (đã nhả khóa) — ghi sổ + cập nhật state ngoài khóa
                        if pos:
                            analytics.record_open(pos)
                            with _lock:
                                _state['positions'].append(pos)
                            _persist_state()
                            open_coins.add(coin)
                            okx_now[coin] = {'inst_id': pos['swap_id']}  # đánh dấu là của ta cho vòng sau
                            _log(f"[{coin}] Mở ✓ giá=${pos['entry_price']:.2f}")
                            notifier.notify_event('OPEN_LONG',
                                f"Mở <b>{coin}</b> ${amount:.0f} @ {opp['funding_rate']*100:.4f}%/8h")
                else:
                    _log("Không lấy được funding rate")
            last_scan = now

        # ── Ghi Excel mỗi 8 giờ (offload, không chặn bot) ───────
        if now - last_excel >= EXCL_INT:
            with _lock:
                ps = [dict(p) for p in _state['positions']]
                u  = _state['usdt']
            if ps:
                def _do_excel(positions=ps, usdt=u):
                    pl = [estimate_pnl(p, funding_override=p.get('_funding_actual')) for p in positions]
                    log_pnl_snapshot(positions, pl, usdt)
                    _log("Ghi Excel ✓ → pnl_log.xlsx")
                _offload('excel', _do_excel)
            last_excel = now

        # ── Push GitHub Pages mỗi 5 phút (offload) ──────────────
        if now - last_push >= PUSH_INT:
            with _lock:
                ps = [dict(p) for p in _state['positions']]
                u  = _state['usdt']
            opps_copy = list(last_opps)
            def _do_push(positions=ps, usdt=u, opps=opps_copy):
                pl = [p.get('_pnl') or estimate_pnl(p, funding_override=p.get('_funding_actual')) for p in positions]
                export_json({'positions': positions, 'pnl_list': pl,
                             'opportunities': opps, 'usdt': usdt})
                if push_to_github():
                    _log("GitHub Pages ✓ (data.json)")
            _offload('push', _do_push)
            last_push = now

        # ── Prune telemetry cũ + checkpoint WAL DB mỗi giờ (chống DB/.db-wal phình) ─
        if now - last_ckpt >= WAL_CKPT_INT:
            try:
                ns, nt = analytics.prune()
                if ns or nt:
                    _log(f"Prune DB: xóa {ns} scan + {nt} tick cũ")
                wal_sz = analytics.wal_checkpoint()
                if wal_sz > 50_000_000:
                    notifier.notify_critical(f"analytics.db-wal vẫn lớn ({wal_sz//1_000_000}MB) sau checkpoint — kiểm tra DB",
                                             key='wal-bloat', dedup_sec=86400)
            except Exception as e:
                _log(f"WAL checkpoint lỗi: {e}")
            last_ckpt = now

        time.sleep(TICK)

    # ── Đóng tất cả khi dừng (dùng race-guard) ──────────────────
    with _lock:
        ps = list(_state['positions'])
    if ps:
        _log(f"Đóng {len(ps)} vị thế...")
        for p in ps:
            _close_one(p, reason='bot_stop')
    _log("━━━ Bot đã dừng ━━━")


def _bot_safe():
    """Wrapper NEVER-DIE: tự restart KHÔNG giới hạn với exp backoff + alert CRITICAL mỗi lần
    crash. Reset đếm nếu chạy ổn định > RESET_AFTER. Chỉ dừng khi user stop (running=False).

    Fix gốc 'sau 3 crash bot chết im': trong 1 tháng không đụng tay, sự cố mạng/OKX có thể
    làm crash >3 lần — bot phải tự sống lại, không được bỏ cuộc."""
    _my_thread = threading.current_thread()
    RESET_AFTER = 300   # chạy ổn > 5 phút → coi là khỏe, reset backoff
    attempt = 0
    while True:
        with _lock:
            still = _state['running'] and _bot_thread is _my_thread
        if not still:
            break
        start = time.time()
        try:
            _bot()
            break   # _bot() return bình thường = user stop → thoát hẳn
        except Exception:
            ran = time.time() - start
            if ran > RESET_AFTER:
                attempt = 0
            attempt += 1
            tb = traceback.format_exc()
            _log(f"━━━ BOT CRASH (lần {attempt}) ━━━")
            for line in tb.splitlines():
                _log(line)
            last_line = (tb.strip().splitlines() or ['?'])[-1]
            notifier.notify_critical(
                f"Bot ARB crash lần {attempt} — tự restart. Lỗi: {last_line[:200]}",
                key='bot-crash', dedup_sec=120)
            with _lock:
                still = _state['running'] and _bot_thread is _my_thread
            if not still:
                break
            delay = min(300, 30 * (2 ** min(attempt - 1, 4)))
            _log(f"Tự restart sau {delay}s...")
            time.sleep(delay)
    with _lock:
        if _bot_thread is _my_thread:
            _state['running'] = False
        _state['closing'].clear()
    try: _ws.stop()
    except Exception: pass
    _log("━━━ Bot kết thúc — state đã reset ━━━")


# ── API ───────────────────────────────────────────────────────────

def _build_status_payload():
    with _lock:
        ps      = list(_state['positions'])
        opps    = list(_state['opportunities'])
        usdt    = _state['usdt']
        teq     = _state['total_eq']
        running = _state['running']
        upd     = _state['last_update']
        logs    = list(_state['log'][-120:])

    out_ps = []
    for p in ps:
        pnl = p.get('_pnl') or {}
        vin = p['contracts'] * p['ct_val'] * p['entry_price']
        pct = (pnl.get('total_pnl', 0) / vin * 100) if vin else 0
        out_ps.append({
            'coin':        p['coin'],
            'entry_price': p['entry_price'],
            'cur_price':   pnl.get('price'),
            'usdt_in':     round(vin, 2),
            'funding_pnl': round(pnl.get('funding_pnl', 0), 4),
            'funding_actual': pnl.get('funding_actual', False),
            'price_pnl':   round(pnl.get('price_pnl', 0), 4),
            'total_pnl':   round(pnl.get('total_pnl', 0), 4),
            'net_pnl':     round(pnl.get('net_pnl', 0), 4),
            'fee_est':     round(pnl.get('fee_est', 0), 4),
            'pct':         round(pct, 3),
            'n_pay':       pnl.get('n_payments', 0),
            'open_time':   datetime.fromtimestamp(p['open_time']).strftime('%d/%m %H:%M'),
            'exit':        p.get('_exit', False),
        })

    return {
        'running': running, 'usdt': usdt, 'total_eq': teq, 'min_rate': MIN_FUNDING_RATE,
        'positions': out_ps,
        'opps':    [{'coin': o['coin'], 'rate': o['funding_rate'],
                     'apy': o['annualized'], 'next': o['next_rate']} for o in opps],
        'logs': logs, 'last_update': upd,
        'ws':   _ws.health(),
        'ts':   time.time(),
    }


@app.route('/api/status')
def api_status():
    return jsonify(_build_status_payload())


@app.route('/api/stream')
def api_stream():
    """Server-Sent Events: push state diffs ~1s for true real-time UI."""
    @stream_with_context
    def gen():
        last = None
        heartbeat = 0
        start = time.time()
        while True:
            # Tự kết thúc sau SSE_TTL (client EventSource sẽ tự reconnect)
            if time.time() - start > SSE_TTL:
                return
            try:
                payload = _build_status_payload()
                data = json.dumps(payload, default=str, separators=(',', ':'))
                if data != last:
                    yield f"event: state\ndata: {data}\n\n"
                    last = data
                    heartbeat = 0
                else:
                    heartbeat += 1
                    if heartbeat >= 15:
                        yield ": keep-alive\n\n"
                        heartbeat = 0
                time.sleep(1.0)
            except GeneratorExit:
                return
            except Exception:
                time.sleep(2.0)
    return Response(
        gen(),
        mimetype='text/event-stream',
        headers={
            'Cache-Control': 'no-cache, no-transform',
            'X-Accel-Buffering': 'no',
            'Connection': 'keep-alive',
        },
    )


def _start_bot():
    """Khởi động bot thread (idempotent). Trả True nếu vừa start, False nếu đang chạy."""
    global _bot_thread
    with _lock:
        if _state['running']:
            return False
        _state['running'] = True
    _bot_thread = threading.Thread(target=_bot_safe, daemon=True, name='bot-main')
    _bot_thread.start()
    return True


@app.route('/api/start', methods=['POST'])
def api_start():
    if not _start_bot():
        return jsonify({'ok': False, 'msg': 'Bot đang chạy rồi'})
    return jsonify({'ok': True})


@app.route('/api/stop', methods=['POST'])
def api_stop():
    with _lock:
        if not _state['running']:
            return jsonify({'ok': False, 'msg': 'Bot chưa chạy'})
        _state['running'] = False
    return jsonify({'ok': True})


@app.route('/api/close/<coin>', methods=['POST'])
def api_close(coin):
    with _lock:
        pos = next((p for p in _state['positions'] if p['coin'] == coin), None)
        already_closing = coin in _state['closing']
    if not pos:
        return jsonify({'ok': False, 'msg': f'Không tìm thấy {coin}'})
    if already_closing:
        return jsonify({'ok': False, 'msg': f'{coin} đang được đóng — chờ hoàn tất'})
    _log(f"[{coin}] Đóng thủ công...")
    threading.Thread(target=_close_one, args=(pos, 'manual'),
                     daemon=True, name=f'close-{coin}').start()
    return jsonify({'ok': True})


@app.route('/api/close_all', methods=['POST'])
def api_close_all():
    """Đóng tất cả vị thế đang mở."""
    with _lock:
        ps = [p for p in _state['positions'] if p['coin'] not in _state['closing']]
    for p in ps:
        threading.Thread(target=_close_one, args=(p, 'manual_all'),
                         daemon=True, name=f"close-{p['coin']}").start()
    return jsonify({'ok': True, 'closed': len(ps)})


@app.route('/api/sync', methods=['POST'])
def api_sync():
    """Ép sync ngay với OKX, xóa phantom positions."""
    try:
        n = _reconcile_with_okx()
        return jsonify({'ok': True, 'removed': n,
                        'msg': f'Đã xóa {n} vị thế phantom' if n else 'Mọi vị thế khớp với OKX'})
    except Exception as e:
        return jsonify({'ok': False, 'msg': str(e)})


@app.route('/api/analytics')
def api_analytics():
    return jsonify({
        'global':  analytics.global_stats(),
        'coins':   analytics.coin_stats(),
        'recent':  analytics.recent_trades(15),
        'equity':  analytics.equity_curve(),
    })


@app.route('/api/pnl-history')
def api_pnl():
    excel_path = os.path.join(os.path.dirname(__file__), EXCEL_FILE)
    if not os.path.exists(excel_path):
        return jsonify({'headers': [], 'rows': []})
    try:
        wb   = openpyxl.load_workbook(excel_path, data_only=True)
        ws   = wb.active
        hdrs = [c.value for c in ws[1] if c.value]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(v is not None and v != '' for v in row):
                rows.append([str(v) if v is not None else '' for v in row[:len(hdrs)]])
        wb.close()
        return jsonify({'headers': hdrs, 'rows': rows})
    except Exception as e:
        return jsonify({'error': str(e), 'headers': [], 'rows': []})


@app.route('/api/health')
def api_health():
    with _lock:
        running = _state['running']
        n_pos   = len(_state['positions'])
    return jsonify({
        'ok':        True,
        'bot':       'okx-arb',
        'running':   running,
        'positions': n_pos,
        'ts':        time.time(),
    })


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/docs')
def docs():
    return render_template('docs.html')


@app.route('/deploy')
def deploy():
    return render_template('deploy.html')


if __name__ == '__main__':
    _port = int(os.getenv('DASH_PORT', '5000'))
    print("\n" + "="*50)
    print("  OKX Arb Bot — Web Dashboard")
    print(f"  Mở trình duyệt: http://localhost:{_port}")
    print("="*50 + "\n")
    # AUTO_START_BOT=true → tự bật bot khi Flask khởi động (không cần curl /api/start).
    # Giúp systemd Restart=always tự hồi phục hoàn toàn mà không phụ thuộc ExecStartPost.
    if os.getenv('AUTO_START_BOT', '').strip().lower() in ('1', 'true', 'yes'):
        if _start_bot():
            print("  AUTO_START_BOT=true → bot đã tự khởi động")
    app.run(host=os.getenv('BIND_HOST', '127.0.0.1'), port=_port, debug=False, use_reloader=False, threaded=True)
